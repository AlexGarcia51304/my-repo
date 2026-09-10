import re
import io
import os
import json
import zipfile
import time
import xml.etree.ElementTree as ET
import pandas as pd
import fitz

from fastapi import FastAPI, File, UploadFile, HTTPException, Form
from fastapi.responses import StreamingResponse
from typing import List, Optional

from presidio_analyzer import (
    AnalyzerEngine,
    PatternRecognizer,
    Pattern,
    RecognizerResult
)

from docx import Document


# App
app = FastAPI(
    title="RISBDC PII Masking API",
    description=(
        "Detects and masks Personally Identifiable Information (PII) from "
        "CRM exports using Microsoft Presidio. "
        "Supported formats: CSV, Excel (.xlsx/.xls), JSON, TXT, PDF, Word (.docx), XML, Markdown (.md). "
        "Includes overlap resolution, SSN validation, Luhn credit-card validation, "
        "configurable filter presets, and consistent per-value pseudonymization "
        "(e.g. <PERSON_1>, <PERSON_2> instead of a flat <PERSON> tag)."
    ),
    version="6.0.0",
)

SUPPORTED_EXTENSIONS = {
    ".csv", ".xlsx", ".xls", ".json", ".txt", ".pdf", ".docx", ".xml", ".md"
}

SUPPORTED_LIST = "`.csv`, `.xlsx`, `.xls`, `.json`, `.txt`, `.pdf`, `.docx`, `.xml`, `.md`"


# Presidio Setup
analyzer = AnalyzerEngine()

# US Street Addresses + PO Boxes
analyzer.registry.add_recognizer(
    PatternRecognizer(
        supported_entity="US_ADDRESS",
        patterns=[
            Pattern(
                name="us_address",
                regex=(
                    r"\b\d+\s+[A-Za-z0-9\s]+\s"
                    r"(?:Street|St|Avenue|Ave|Road|Rd|Drive|Dr|Lane|Ln"
                    r"|Boulevard|Blvd|Court|Ct|Circle|Cir|Way|Parkway|Pkwy"
                    r"|Place|Pl|Terrace|Ter|Highway|Hwy)\b"
                ),
                score=1.0,
            ),
            Pattern(
                name="po_box",
                regex=r"\bP\.?O\.?\s*Box\s+\d+\b",
                score=1.0,
            ),
        ],
    )
)

# EIN
analyzer.registry.add_recognizer(
    PatternRecognizer(
        supported_entity="EIN",
        patterns=[Pattern(name="ein", regex=r"\b\d{2}-\d{7}\b", score=1.0)],
    )
)

# SBA Loan Numbers
analyzer.registry.add_recognizer(
    PatternRecognizer(
        supported_entity="SBA_LOAN_NUMBER",
        patterns=[Pattern(name="sba_loan", regex=r"\bSBA-\d{4}-\d+\b", score=1.0)],
    )
)

# Client IDs
analyzer.registry.add_recognizer(
    PatternRecognizer(
        supported_entity="CLIENT_ID",
        patterns=[Pattern(name="client_id", regex=r"\bCLI-\d+\b", score=1.0)],
    )
)

# ZIP Codes
analyzer.registry.add_recognizer(
    PatternRecognizer(
        supported_entity="ZIP_CODE",
        patterns=[Pattern(name="zip_code", regex=r"\b\d{5}(?:-\d{4})?\b", score=1.0)],
    )
)

# Constants
RI_CITIES = [
    "Providence", "Warwick", "Cranston", "Newport", "Pawtucket",
    "Woonsocket", "Coventry", "Cumberland", "Johnston", "North Providence",
    "West Warwick", "North Kingstown", "South Kingstown", "Westerly",
    "Smithfield", "Lincoln", "Portsmouth", "Barrington", "Middletown",
    "Bristol", "Tiverton", "Narragansett", "East Greenwich", "Hopkinton",
    "Scituate", "Glocester", "Charlestown", "Richmond", "Exeter",
    "Little Compton", "Burrillville", "Central Falls", "Jamestown",
    "New Shoreham", "Foster",
]

us_phone_pattern = re.compile(r"\b\d{3}-\d{3}-\d{4}\b")
ssn_pattern = re.compile(r"\b\d{3}-\d{2}-\d{4}\b")
credit_card_pattern = re.compile(r"\b(?:\d[ -]*?){13,16}\b")
address_force_pattern = re.compile(
    r"\b\d+\s+[A-Za-z0-9\s]+\s"
    r"(?:Street|St|Avenue|Ave|Road|Rd|Drive|Dr|Lane|Ln"
    r"|Boulevard|Blvd|Court|Ct|Circle|Cir|Way|Parkway|Pkwy"
    r"|Place|Pl|Terrace|Ter|Highway|Hwy)\b",
    flags=re.IGNORECASE,
)


def is_valid_ssn(ssn_str: str) -> bool:
    """Validate a US SSN using the standard invalid-area/group/serial rules."""
    area, group, serial = ssn_str.split("-")
    if area in ("000", "666") or area.startswith("9"):
        return False
    if group == "00" or serial == "0000":
        return False
    return True


def luhn_check(number_str: str) -> bool:
    """Return True when a candidate card number passes the Luhn checksum."""
    digits = [int(d) for d in re.sub(r"[ -]", "", number_str)]
    if not (13 <= len(digits) <= 19):
        return False

    checksum = 0
    parity = len(digits) % 2
    for i, digit in enumerate(digits):
        if i % 2 == parity:
            digit *= 2
            if digit > 9:
                digit -= 9
        checksum += digit

    return checksum % 10 == 0


def resolve_overlaps(results):
    """
    Resolve overlapping Presidio/custom detections before masking.

    The highest-scoring result wins; if scores tie, the longer span wins.
    This prevents the same PII span from being counted or masked twice.
    """
    if not results:
        return []

    sorted_results = sorted(
        results,
        key=lambda r: (r.start, -r.score, -(r.end - r.start))
    )

    resolved = [sorted_results[0]]

    for result in sorted_results[1:]:
        previous = resolved[-1]

        if result.start >= previous.end:
            resolved.append(result)
        elif (result.score, result.end - result.start) > (
            previous.score,
            previous.end - previous.start,
        ):
            resolved[-1] = result

    return resolved


# Configurable filter groups and/or presets
DEFAULT_PRESETS = {
    "all": None,  
    "financial": [
        "US_SSN", "US_ITIN", "CREDIT_CARD", "EIN", "SBA_LOAN_NUMBER"
    ],
    "contact": [
        "PHONE_NUMBER", "US_ADDRESS", "LOCATION", "ZIP_CODE", "EMAIL_ADDRESS"
    ],
    "identity": [
        "PERSON", "CLIENT_ID"
    ],
}

PRESETS_FILE = os.environ.get("PRESETS_FILE", "presets.json")


def load_presets() -> dict:
    presets = dict(DEFAULT_PRESETS)

    if PRESETS_FILE and os.path.isfile(PRESETS_FILE):
        with open(PRESETS_FILE, "r", encoding="utf-8") as f:
            custom = json.load(f)
        for name, entities in custom.items():
            presets[name] = entities

    return presets


PRESETS = load_presets()


def resolve_preset(name: str) -> Optional[set]:
    """Look up a preset by name and return its active entity set (or None for 'all')."""
    if name not in PRESETS:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Unknown preset '{name}'. "
                f"Available presets: {', '.join(PRESETS.keys())}"
            ),
        )
    entities = PRESETS[name]
    return set(entities) if entities is not None else None


# Shared detection
def detect_entities(text: str, pii_counts: dict, active_entities: Optional[set]):
    text = str(text)

    results = [
        r for r in analyzer.analyze(text=text, language="en")
        if not (r.entity_type == "PERSON" and r.score < 0.75)
    ]

    # Force Rhode Island cities
    for city in RI_CITIES:
        for match in re.finditer(
            rf"\b{re.escape(city)}\b", text, flags=re.IGNORECASE
        ):
            results.append(
                RecognizerResult("LOCATION", match.start(), match.end(), 1.0)
            )

    # Force addresses
    for match in address_force_pattern.finditer(text):
        results.append(
            RecognizerResult("US_ADDRESS", match.start(), match.end(), 1.0)
        )

    # Force phones
    for match in us_phone_pattern.finditer(text):
        results.append(
            RecognizerResult("PHONE_NUMBER", match.start(), match.end(), 1.0)
        )

    # Force SSNs only when structurally valid
    for match in ssn_pattern.finditer(text):
        if is_valid_ssn(match.group()):
            results.append(
                RecognizerResult("US_SSN", match.start(), match.end(), 1.0)
            )

    # Force credit cards only when they pass Luhn validation
    for match in credit_card_pattern.finditer(text):
        if luhn_check(match.group()):
            results.append(
                RecognizerResult("CREDIT_CARD", match.start(), match.end(), 1.0)
            )

    # Resolve overlapping Presidio/custom detections before counting/masking 
    results = resolve_overlaps(results)

    # Apply the active preset (if any) to filter out unwanted entity types 
    if active_entities is not None:
        results = [r for r in results if r.entity_type in active_entities]

    for result in results:
        pii_counts[result.entity_type] = (
            pii_counts.get(result.entity_type, 0) + 1
        )

    return results


# Consistent PII Tags
TAG_ALIAS = {
    "US_ITIN": "US_SSN",
}


def assign_tag(entity_type: str, value: str, value_map: dict, type_counters: dict) -> str:
    tag_type = TAG_ALIAS.get(entity_type, entity_type)
    key = (tag_type, value.strip().lower())

    if key in value_map:
        return value_map[key]

    type_counters[tag_type] = type_counters.get(tag_type, 0) + 1
    tag = f"<{tag_type}_{type_counters[tag_type]}>"
    value_map[key] = tag
    return tag


def render_masked_text(text: str, results, value_map: dict, type_counters: dict) -> str:
    if not results:
        return text

    pieces = []
    pos = 0

    for result in sorted(results, key=lambda r: r.start):
        if result.start > pos:
            pieces.append(text[pos:result.start])

        original_value = text[result.start:result.end]
        pieces.append(
            assign_tag(result.entity_type, original_value, value_map, type_counters)
        )
        pos = result.end

    pieces.append(text[pos:])
    return "".join(pieces)


# Core Masking Logic

def mask_pii(
    text: str,
    pii_counts: dict,
    value_map: dict,
    type_counters: dict,
    active_entities: Optional[set],
) -> str:
    """Mask PII in a single text value. Updates pii_counts in-place."""
    text = str(text)
    results = detect_entities(text, pii_counts, active_entities)
    return render_masked_text(text, results, value_map, type_counters)


def mask_json_obj(obj, pii_counts: dict, value_map: dict, type_counters: dict, active_entities: Optional[set]):
    """Recursively mask PII in a parsed JSON object."""
    if isinstance(obj, dict):
        return {
            k: mask_json_obj(v, pii_counts, value_map, type_counters, active_entities)
            for k, v in obj.items()
        }
    elif isinstance(obj, list):
        return [
            mask_json_obj(i, pii_counts, value_map, type_counters, active_entities)
            for i in obj
        ]
    elif isinstance(obj, str):
        return mask_pii(obj, pii_counts, value_map, type_counters, active_entities)
    else:
        return obj


def mask_xml_element(element, pii_counts: dict, value_map: dict, type_counters: dict, active_entities: Optional[set]):
    """Recursively mask PII in XML text, tail text, and attribute values."""
    if element.text:
        element.text = mask_pii(element.text, pii_counts, value_map, type_counters, active_entities)

    if element.tail:
        element.tail = mask_pii(element.tail, pii_counts, value_map, type_counters, active_entities)

    for attr_name, attr_value in element.attrib.items():
        element.set(
            attr_name,
            mask_pii(attr_value, pii_counts, value_map, type_counters, active_entities),
        )

    for child in element:
        mask_xml_element(child, pii_counts, value_map, type_counters, active_entities)


def process_dataframe(
    df: pd.DataFrame,
    pii_counts: dict,
    value_map: dict,
    type_counters: dict,
    active_entities: Optional[set],
) -> pd.DataFrame:
    """Mask every cell in a DataFrame. Updates pii_counts."""
    masked_df = df.copy()
    for col in masked_df.columns:
        masked_df[col] = masked_df[col].astype(str).apply(
            lambda v: mask_pii(v, pii_counts, value_map, type_counters, active_entities)
        )
    return masked_df


# Multi-span entity detection for docx
def mask_paragraph(paragraph, pii_counts: dict, value_map: dict, type_counters: dict, active_entities: Optional[set]):
    runs = paragraph.runs

    if not runs:
        if paragraph.text:
            paragraph.text = mask_pii(
                paragraph.text, pii_counts, value_map, type_counters, active_entities
            )
        return

    full_text = "".join(run.text for run in runs)

    if not full_text:
        return

    results = detect_entities(full_text, pii_counts, active_entities)

    if not results:
        return

    run_spans = []
    cursor = 0
    for run in runs:
        run_spans.append((cursor, cursor + len(run.text)))
        cursor += len(run.text)

    for run, (run_start, run_end) in zip(runs, run_spans):
        pos = run_start
        segments = []

        for result in results:
            if result.end <= run_start or result.start >= run_end:
                continue 

            seg_start = max(result.start, pos)
            if seg_start > pos:
                segments.append(full_text[pos:seg_start])

            if result.start >= run_start:
                tag = assign_tag(
                    result.entity_type,
                    full_text[result.start:result.end],
                    value_map,
                    type_counters,
                )
                segments.append(tag)

            pos = min(result.end, run_end)

        if pos < run_end:
            segments.append(full_text[pos:run_end])

        run.text = "".join(segments)


def mask_table(table, pii_counts: dict, value_map: dict, type_counters: dict, active_entities: Optional[set]):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                mask_paragraph(paragraph, pii_counts, value_map, type_counters, active_entities)
            for nested_table in cell.tables:
                mask_table(nested_table, pii_counts, value_map, type_counters, active_entities)


def mask_region(region, pii_counts: dict, value_map: dict, type_counters: dict, active_entities: Optional[set]):
    for paragraph in region.paragraphs:
        mask_paragraph(paragraph, pii_counts, value_map, type_counters, active_entities)
    for table in region.tables:
        mask_table(table, pii_counts, value_map, type_counters, active_entities)


def calculate_risk_level(pii_counts: dict) -> str:
    """Calculate file risk using the upgraded prototype thresholds."""
    total = sum(pii_counts.values())

    if pii_counts.get("US_SSN", 0) or pii_counts.get("CREDIT_CARD", 0):
        return "HIGH"
    if total >= 30:
        return "HIGH"
    if total >= 10:
        return "MEDIUM"
    return "LOW"


def build_report(
    pii_counts: dict,
    processing_time_seconds: float | None = None,
    preset: Optional[str] = None,
) -> dict:
    """Build the structured PII report payload."""
    total = sum(pii_counts.values())
    breakdown = {
        entity: {
            "count": count,
            "percentage": round(count / total * 100, 1) if total else 0.0,
        }
        for entity, count in sorted(
            pii_counts.items(), key=lambda x: x[1], reverse=True
        )
    }

    report = {
        "total_pii_masked": total,
        "breakdown": breakdown,
        "risk_level": calculate_risk_level(pii_counts),
    }

    if preset is not None:
        report["filter_preset"] = preset

    if processing_time_seconds is not None:
        report["processing_time_seconds"] = round(processing_time_seconds, 2)

    return report

def get_extension(filename: str) -> str:
    return "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

def validate_extension(filename: str) -> str:
    ext = get_extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=(
                f"Unsupported file type '{ext}'. "
                f"Accepted: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
            ),
        )
    return ext

def masked_filename_for(original: str) -> str:
    """
    Return the output filename for a masked file.
    .xls is always saved back as .xlsx.
    PDF stays .pdf (redacted in-place via fitz).
    """
    ext = get_extension(original)
    stem = original.rsplit(".", 1)[0] if "." in original else original

    if ext == ".xls":
        return f"masked_{stem}.xlsx"
    else:
        return f"masked_{original}"

def process_raw(
    ext: str,
    raw: bytes,
    pii_counts: dict,
    value_map: dict,
    type_counters: dict,
    active_entities: Optional[set],
) -> tuple[bytes, str]:
    """
    Mask raw file bytes for a given extension.
    Returns (masked_bytes, media_type).
    """

    # CSV
    if ext == ".csv":
        df = pd.read_csv(
            io.StringIO(raw.decode("utf-8")), dtype=str, keep_default_na=False
        )
        masked_df = process_dataframe(df, pii_counts, value_map, type_counters, active_entities)
        buf = io.StringIO()
        masked_df.to_csv(buf, index=False)
        return buf.getvalue().encode("utf-8"), "text/csv"

    # Excel (.xlsx or .xls , always saved back as .xlsx)
    elif ext in (".xlsx", ".xls"):
        sheets = pd.read_excel(io.BytesIO(raw), sheet_name=None, dtype=str)
        out_buf = io.BytesIO()
        with pd.ExcelWriter(out_buf, engine="openpyxl") as writer:
            for sheet_name, df in sheets.items():
                df = df.fillna("")
                masked_df = process_dataframe(df, pii_counts, value_map, type_counters, active_entities)
                masked_df.to_excel(writer, sheet_name=sheet_name, index=False)
        return (
            out_buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    # JSON
    elif ext == ".json":
        data = json.loads(raw.decode("utf-8"))
        masked = mask_json_obj(data, pii_counts, value_map, type_counters, active_entities)
        return json.dumps(masked, indent=4).encode("utf-8"), "application/json"

    # TXT
    elif ext == ".txt":
        masked = mask_pii(raw.decode("utf-8"), pii_counts, value_map, type_counters, active_entities)
        return masked.encode("utf-8"), "text/plain"

    # Markdown
    elif ext == ".md":
        masked = mask_pii(raw.decode("utf-8"), pii_counts, value_map, type_counters, active_entities)
        return masked.encode("utf-8"), "text/markdown"

    # PDF , redacted instead of masked
    elif ext == ".pdf":
        pdf = fitz.open(stream=raw, filetype="pdf")
        for page in pdf:
            blocks = page.get_text("dict")["blocks"]
            for block in blocks:
                if "lines" not in block:
                    continue

                # Full-block PII detection across text spans

                span_list = [
                    span
                    for line in block["lines"]
                    for span in line["spans"]
                ]

                if not span_list:
                    continue

                block_text = "".join(span["text"] for span in span_list)

                if not block_text.strip():
                    continue

                results = detect_entities(block_text, pii_counts, active_entities)

                if not results:
                    continue

                span_spans = []
                cursor = 0
                for span in span_list:
                    span_spans.append((cursor, cursor + len(span["text"])))
                    cursor += len(span["text"])

                for span, (span_start, span_end) in zip(span_list, span_spans):
                    touched = any(
                        r.start < span_end and r.end > span_start
                        for r in results
                    )
                    if touched:
                        rect = fitz.Rect(span["bbox"])
                        page.add_redact_annot(rect, fill=(0, 0, 0))

            page.apply_redactions()
        out_buf = io.BytesIO()
        pdf.save(out_buf)
        pdf.close()
        return out_buf.getvalue(), "application/pdf"

    # Word (.docx)
    elif ext == ".docx":
        doc = Document(io.BytesIO(raw))

        # Main document paragraphs and tables
        for paragraph in doc.paragraphs:
            mask_paragraph(paragraph, pii_counts, value_map, type_counters, active_entities)

        for table in doc.tables:
            mask_table(table, pii_counts, value_map, type_counters, active_entities)

        # Headers and footers, including their tables
        for section in doc.sections:
            regions = (
                section.header,
                section.footer,
                section.first_page_header,
                section.first_page_footer,
                section.even_page_header,
                section.even_page_footer,
            )
            for region in regions:
                mask_region(region, pii_counts, value_map, type_counters, active_entities)

        out_buf = io.BytesIO()
        doc.save(out_buf)
        return (
            out_buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    # XML
    elif ext == ".xml":
        root = ET.fromstring(raw.decode("utf-8"))
        mask_xml_element(root, pii_counts, value_map, type_counters, active_entities)
        out_buf = io.BytesIO()
        tree = ET.ElementTree(root)
        tree.write(out_buf, encoding="utf-8", xml_declaration=True)
        return out_buf.getvalue(), "application/xml"

    else:
        raise ValueError(f"Unhandled extension: {ext}")

async def mask_upload(file: UploadFile, preset: str = "all") -> tuple[bytes, str, dict, str, float, str]:
    """
    Read an uploaded file, mask it, and return:
      (masked_bytes, masked_filename, pii_counts, media_type, processing_time_seconds, preset)
    """
    start_time = time.perf_counter()

    ext = validate_extension(file.filename)
    active_entities = resolve_preset(preset)

    raw = await file.read()
    pii_counts: dict = {}
    value_map: dict = {}
    type_counters: dict = {}

    try:
        masked_bytes, media_type = process_raw(
            ext, raw, pii_counts, value_map, type_counters, active_entities
        )
    except Exception as e:
        raise HTTPException(
            status_code=422,
            detail=f"Failed to process '{file.filename}': {e}",
        )

    processing_time = time.perf_counter() - start_time

    return (
        masked_bytes,
        masked_filename_for(file.filename),
        pii_counts,
        media_type,
        processing_time,
        preset,
    )


# Endpoints

@app.get("/", summary="Health check")
def root():
    """Confirms the API is running."""
    return {"status": "ok", "service": "RISBDC PII Masking API", "version": "6.0.0"}


@app.get(
    "/presets",
    summary="List available filter presets",
)
def list_presets():
    """
    Returns every available filter preset and the entity types it masks.
    "all" masks every entity type Presidio and the custom recognizers detect;
    every other preset masks only the entity types listed.

    Add your own presets without touching the code by placing a JSON file
    (default path: `presets.json`, override with the `PRESETS_FILE` env var)
    alongside the service, e.g. `{"my_preset": ["PERSON", "PHONE_NUMBER"]}`.
    """
    return {
        name: ("all detected entity types" if entities is None else entities)
        for name, entities in PRESETS.items()
    }


@app.post(
    "/mask",
    summary="Mask a single file — download masked file directly",
    response_class=StreamingResponse,
    responses={
        200: {"description": "Masked file (same format as input)"},
        400: {"description": "Unsupported file type or unknown preset"},
        422: {"description": "File could not be processed"},
    },
)
async def mask_single(
    file: UploadFile = File(
        ...,
        description=f"File to mask. Supported: {SUPPORTED_LIST}",
    ),
    preset: str = Form(
        "all",
        description="Filter preset controlling which entity types get masked. See GET /presets for options.",
    ),
):
    """
    Upload a single file and receive the masked version as a direct download.

    Detected PII is replaced with a consistent, numbered tag per unique value
    (e.g. `<PERSON_1>`, `<PERSON_2>`) — the same value is tagged identically
    everywhere it appears in the file, so different people/addresses/etc.
    stay distinguishable from each other in the masked output. SSNs are
    validated before forced detection, credit cards are validated with Luhn,
    and overlapping detections are resolved before masking.

    **Note on PDF:** PII is redacted directly inside the PDF using black rectangles —
    the output is a proper `.pdf` file with PII blacked out, not a text extraction.

    **Note on .xls:** Legacy Excel files are always saved back as `.xlsx`.
    """
    masked_bytes, out_filename, _, media_type, _, _ = await mask_upload(file, preset)
    return StreamingResponse(
        iter([masked_bytes]),
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{out_filename}"'},
    )

@app.post(
    "/mask/report",
    summary="Mask a single file — download masked file + PII report together",
    response_class=StreamingResponse,
    responses={
        200: {
            "content": {"application/zip": {}},
            "description": "ZIP containing the masked file and pii_report.json",
        },
        400: {"description": "Unsupported file type or unknown preset"},
        422: {"description": "File could not be processed"},
    },
)
async def mask_with_report(
    file: UploadFile = File(
        ...,
        description=f"File to mask. Supported: {SUPPORTED_LIST}",
    ),
    preset: str = Form(
        "all",
        description="Filter preset controlling which entity types get masked. See GET /presets for options.",
    ),
):
    """
    Upload a single file. Returns a ZIP archive containing:
    - `masked_<filename>` — the fully masked file
    - `pii_report.json` — PII detection breakdown, risk level, filter preset, and processing time

    Both download together in one click.
    """
    masked_bytes, out_filename, pii_counts, _, processing_time, used_preset = await mask_upload(file, preset)

    report_payload = {
        "filename": file.filename,
        "report": build_report(pii_counts, processing_time, used_preset),
    }

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(out_filename, masked_bytes)
        zf.writestr("pii_report.json", json.dumps(report_payload, indent=2))
    zip_buffer.seek(0)

    stem = file.filename.rsplit(".", 1)[0] if "." in file.filename else file.filename
    return StreamingResponse(
        iter([zip_buffer.getvalue()]),
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="masked_{stem}_results.zip"'},
    )

@app.post(
    "/report",
    summary="Report only — scan for PII without masking",
)
async def report_only(
    file: UploadFile = File(
        ...,
        description=f"File to scan. Supported: {SUPPORTED_LIST}",
    ),
    preset: str = Form(
        "all",
        description="Filter preset controlling which entity types are scanned for. See GET /presets for options.",
    ),
):
    """
    Scan a file for PII and return only the detection report (no masked output).
    Useful for auditing before committing to full masking.
    """
    _, _, pii_counts, _, processing_time, used_preset = await mask_upload(file, preset)
    return {"filename": file.filename, "report": build_report(pii_counts, processing_time, used_preset)}