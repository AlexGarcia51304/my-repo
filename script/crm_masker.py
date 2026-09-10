import re
import glob
import json
import argparse
import xml.etree.ElementTree as ET
from pathlib import Path
import pandas as pd
from docx import Document
import fitz
import os
import time
from datetime import datetime

from presidio_analyzer import (
  AnalyzerEngine,
  PatternRecognizer,
  Pattern,
  RecognizerResult
)


# Initialize Presidio
analyzer = AnalyzerEngine()

# PII Detection Report
pii_counts = {}

# Stores PII counts for each file
file_pii_counts = {}
file_risk = {}

### Custom Recognizers
# US Street Addresses
address_pattern = Pattern(
  name="us_address",
  regex=r"\b\d+\s+[A-Za-z0-9\s]+\s(?:Street|St|Avenue|Ave|Road|Rd|Drive|Dr|Lane|Ln|Boulevard|Blvd|Court|Ct|Circle|Cir|Way|Parkway|Pkwy|Place|Pl|Terrace|Ter|Highway|Hwy)\b",
  score=1.0
)

analyzer.registry.add_recognizer(
  PatternRecognizer(
      supported_entity="US_ADDRESS",
      patterns=[address_pattern]
  )
)

# PO Boxes
po_box_pattern = Pattern(
  name="po_box",
  regex=r"\bP\.?O\.?\s*Box\s+\d+\b",
  score=1.0
)

analyzer.registry.add_recognizer(
  PatternRecognizer(
      supported_entity="US_ADDRESS",
      patterns=[po_box_pattern]
  )
)

# EIN
ein_pattern = Pattern(
  name="ein",
  regex=r"\b\d{2}-\d{7}\b",
  score=1.0
)

analyzer.registry.add_recognizer(
  PatternRecognizer(
      supported_entity="EIN",
      patterns=[ein_pattern]
  )
)

# SBA Loan Numbers
sba_pattern = Pattern(
  name="sba_loan",
  regex=r"\bSBA-\d{4}-\d+\b",
  score=1.0
)

analyzer.registry.add_recognizer(
  PatternRecognizer(
      supported_entity="SBA_LOAN_NUMBER",
      patterns=[sba_pattern]
  )
)

# Client IDs
client_pattern = Pattern(
  name="client_id",
  regex=r"\bCLI-\d+\b",
  score=1.0
)

analyzer.registry.add_recognizer(
  PatternRecognizer(
      supported_entity="CLIENT_ID",
      patterns=[client_pattern]
  )
)

# ZIP Codes
zip_pattern = Pattern(
  name="zip_code",
  regex=r"\b\d{5}(?:-\d{4})?\b",
  score=1.0
)

analyzer.registry.add_recognizer(
  PatternRecognizer(
      supported_entity="ZIP_CODE",
      patterns=[zip_pattern]
  )
)

# Rhode Island Cities
RI_CITIES = [
  "Providence",
  "Warwick",
  "Cranston",
  "Newport",
  "Pawtucket",
  "Woonsocket",
  "Coventry",
  "Cumberland",
  "Johnston",
  "North Providence",
  "West Warwick",
  "North Kingstown",
  "South Kingstown",
  "Westerly",
  "Smithfield",
  "Lincoln",
  "Portsmouth",
  "Barrington",
  "Middletown",
  "Bristol",
  "Tiverton",
  "Narragansett",
  "East Greenwich",
  "Hopkinton",
  "Scituate",
  "Glocester",
  "Charlestown",
  "Richmond",
  "Exeter",
  "Little Compton",
  "Burrillville",
  "Central Falls",
  "Jamestown",
  "New Shoreham",
  "Foster"
]

def resolve_overlaps(results):
    """
    Resolve overlapping RecognizerResult spans before anonymizing.
    Keeps the highest-scoring (and longest, on ties) result for
    any given stretch of text so the same PII isn't double-counted
    or double-masked under two different entity types.
    """
    if not results:
        return []

    sorted_results = sorted(
        results,
        key=lambda r: (r.start, -r.score, -(r.end - r.start))
    )

    resolved = [sorted_results[0]]

    for r in sorted_results[1:]:
        prev = resolved[-1]
        if r.start >= prev.end:
            resolved.append(r)
        elif (r.score, r.end - r.start) > (prev.score, prev.end - prev.start):
            resolved[-1] = r

    return resolved

# Validation Patterns
us_phone_pattern = re.compile(
  r"\b\d{3}-\d{3}-\d{4}\b"
)

ssn_pattern = re.compile(
  r"\b\d{3}-\d{2}-\d{4}\b"
)

def is_valid_ssn(ssn_str):
    area, group, serial = ssn_str.split("-")
    if area in ("000", "666") or area.startswith("9"):
        return False
    if group == "00" or serial == "0000":
        return False
    return True

credit_card_pattern = re.compile(
  r"\b(?:\d[ -]*?){13,16}\b"
)

def luhn_check(number_str):
    digits = [int(d) for d in re.sub(r"[ -]", "", number_str)]
    if not (13 <= len(digits) <= 19):
        return False
    checksum = 0
    parity = len(digits) % 2
    for i, d in enumerate(digits):
        if i % 2 == parity:
            d *= 2
            if d > 9:
                d -= 9
        checksum += d
    return checksum % 10 == 0

address_force_pattern = re.compile(
  r"\b\d+\s+[A-Za-z0-9\s]+\s(?:Street|St|Avenue|Ave|Road|Rd|Drive|Dr|Lane|Ln|Boulevard|Blvd|Court|Ct|Circle|Cir|Way|Parkway|Pkwy|Place|Pl|Terrace|Ter|Highway|Hwy)\b",
  flags=re.IGNORECASE
)


# Configurable filter groups and/or presents
DEFAULT_PRESETS = {
    "all": None,  
    "financial": [
        "US_SSN", "US_ITIN", "CREDIT_CARD", "EIN", "SBA_LOAN_NUMBER"
    ],
    "contact": [
        "PHONE_NUMBER", "US_ADDRESS", "LOCATION", "ZIP_CODE", "EMAIL_ADDRESS"
    ],
    "identity": [
        "PERSON", "CLIENT_ID"
    ],
}


def load_presets(presets_file):
    """
    Start from the built-in defaults, then let a presets.json file (if
    present) add new presets or override existing ones by name.
    """
    presets = dict(DEFAULT_PRESETS)

    if presets_file and os.path.isfile(presets_file):
        with open(presets_file, "r", encoding="utf-8") as f:
            custom = json.load(f)

        for name, entities in custom.items():
            presets[name] = entities

    return presets


def parse_args():
    parser = argparse.ArgumentParser(
        description="RISBDC CRM PII Masking Tool"
    )
    parser.add_argument(
        "--preset",
        default="all",
        help="Name of the filter preset/group to apply (default: all)."
    )
    parser.add_argument(
        "--presets-file",
        default="presets.json",
        help=(
            "Path to a JSON file defining custom presets, e.g. "
            '{"my_preset": ["PERSON", "US_SSN"]}. '
            "Defaults to presets.json in the current directory if present."
        )
    )
    parser.add_argument(
        "--list-presets",
        action="store_true",
        help="List available presets and exit."
    )
    return parser.parse_args()


args = parse_args()
PRESETS = load_presets(args.presets_file)

if args.list_presets:
    print("Available presets:\n")
    for name, entities in PRESETS.items():
        scope = (
            "ALL detected entity types"
            if entities is None
            else ", ".join(entities)
        )
        print(f"  {name}: {scope}")
    raise SystemExit(0)

if args.preset not in PRESETS:
    raise SystemExit(
        f"Unknown preset '{args.preset}'. "
        f"Available presets: {', '.join(PRESETS.keys())}"
    )

ACTIVE_ENTITIES = PRESETS[args.preset]
if ACTIVE_ENTITIES is not None:
    ACTIVE_ENTITIES = set(ACTIVE_ENTITIES)

print(
    f"Using filter preset: {args.preset}"
    + (
        f" ({', '.join(sorted(ACTIVE_ENTITIES))})"
        if ACTIVE_ENTITIES
        else " (all detected entity types)"
    )
)


# Shared detection
def detect_entities(text, file_counts=None):

    global pii_counts

    text = str(text)

    results = [
        r for r in analyzer.analyze(
            text=text,
            language="en"
        )
        if not (
            r.entity_type == "PERSON"
            and r.score < 0.75
        )
    ]

    # Force Rhode Island Cities
    for city in RI_CITIES:
        for match in re.finditer(
            rf"\b{re.escape(city)}\b",
            text,
            flags=re.IGNORECASE
        ):
            results.append(
                RecognizerResult(
                    entity_type="LOCATION",
                    start=match.start(),
                    end=match.end(),
                    score=1.0
                )
            )

    # Force Addresses
    for match in address_force_pattern.finditer(text):
        results.append(
            RecognizerResult(
                entity_type="US_ADDRESS",
                start=match.start(),
                end=match.end(),
                score=1.0
            )
        )

    # Force Phones
    for match in us_phone_pattern.finditer(text):
        results.append(
            RecognizerResult(
                entity_type="PHONE_NUMBER",
                start=match.start(),
                end=match.end(),
                score=1.0
            )
        )

    # Force SSNs
    for match in ssn_pattern.finditer(text):
        if is_valid_ssn(match.group()):
            results.append(
                RecognizerResult(
                    entity_type="US_SSN",
                    start=match.start(),
                    end=match.end(),
                    score=1.0
                )
            )

    # Force Credit Cards
    for match in credit_card_pattern.finditer(text):
        if luhn_check(match.group()):
            results.append(
                RecognizerResult(
                    entity_type="CREDIT_CARD",
                    start=match.start(),
                    end=match.end(),
                    score=1.0
                )
            )

    results = resolve_overlaps(results)

    # Apply the active preset 
    if ACTIVE_ENTITIES is not None:
        results = [r for r in results if r.entity_type in ACTIVE_ENTITIES]

    # Count detections
    for result in results:
        entity = result.entity_type

        pii_counts[entity] = pii_counts.get(entity, 0) + 1

        if file_counts is not None:
            file_counts[entity] = file_counts.get(entity, 0) + 1

    return results


# Consistent PII Tags
TAG_ALIAS = {
    "US_ITIN": "US_SSN",
}


def assign_tag(entity_type, value, value_map, type_counters):
    tag_type = TAG_ALIAS.get(entity_type, entity_type)
    key = (tag_type, value.strip().lower())

    if key in value_map:
        return value_map[key]

    type_counters[tag_type] = type_counters.get(tag_type, 0) + 1
    tag = f"<{tag_type}_{type_counters[tag_type]}>"
    value_map[key] = tag
    return tag


def render_masked_text(text, results, value_map, type_counters):
    if not results:
        return text

    pieces = []
    pos = 0

    for result in sorted(results, key=lambda r: r.start):
        if result.start > pos:
            pieces.append(text[pos:result.start])

        original_value = text[result.start:result.end]
        pieces.append(
            assign_tag(result.entity_type, original_value, value_map, type_counters)
        )
        pos = result.end

    pieces.append(text[pos:])
    return "".join(pieces)


def mask_pii(text, file_counts=None, value_map=None, type_counters=None):
    """
    Detect + mask a single string. Used for CSV/Excel cells, JSON string
    values, and plain text/markdown files, where there's no run/span
    structure to preserve.
    """
    if value_map is None:
        value_map = {}
    if type_counters is None:
        type_counters = {}

    text = str(text)
    results = detect_entities(text, file_counts)
    return render_masked_text(text, results, value_map, type_counters)


# JSON Masking Helper
def mask_json(obj, file_counts=None, value_map=None, type_counters=None):

    if value_map is None:
        value_map = {}
    if type_counters is None:
        type_counters = {}

    if isinstance(obj, dict):
        return {
            k: mask_json(v, file_counts, value_map, type_counters)
            for k, v in obj.items()
        }

    elif isinstance(obj, list):
        return [
            mask_json(i, file_counts, value_map, type_counters)
            for i in obj
        ]

    elif isinstance(obj, str):
        return mask_pii(obj, file_counts, value_map, type_counters)

    else:
        return obj


# Multi-span entity detection for docx
def mask_paragraph(paragraph, file_counts, value_map, type_counters):

    runs = paragraph.runs

    if not runs:
        if paragraph.text:
            paragraph.text = mask_pii(paragraph.text, file_counts, value_map, type_counters)
        return

    full_text = "".join(run.text for run in runs)

    if not full_text:
        return

    results = detect_entities(full_text, file_counts)

    if not results:
        return

    run_spans = []
    cursor = 0
    for run in runs:
        run_spans.append((cursor, cursor + len(run.text)))
        cursor += len(run.text)

    for run, (run_start, run_end) in zip(runs, run_spans):
        pos = run_start
        segments = []

        for result in results:
            if result.end <= run_start or result.start >= run_end:
                continue  

            seg_start = max(result.start, pos)
            if seg_start > pos:
                segments.append(full_text[pos:seg_start])

            if result.start >= run_start:
                tag = assign_tag(
                    result.entity_type,
                    full_text[result.start:result.end],
                    value_map,
                    type_counters
                )
                segments.append(tag)

            pos = min(result.end, run_end)

        if pos < run_end:
            segments.append(full_text[pos:run_end])

        run.text = "".join(segments)


def mask_table(table, file_counts, value_map, type_counters):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                mask_paragraph(p, file_counts, value_map, type_counters)
            for nested_table in cell.tables:
                mask_table(nested_table, file_counts, value_map, type_counters)


# Process Supported Files
SUPPORTED_EXTENSIONS = {
   ".csv",
   ".xlsx",
   ".xls",
   ".json",
   ".txt",
   ".pdf",
   ".docx",
   ".xml",
   ".md"
}

OUTPUT_DIR = "Masked_Output"

os.makedirs(
    OUTPUT_DIR,
    exist_ok=True
)

supported_files = []
unsupported_files = []

for file in glob.glob("*"):

   extension = Path(file).suffix.lower()

   if extension in SUPPORTED_EXTENSIONS:

       supported_files.append(file)

   elif extension != "" and file != Path(__file__).name:

       unsupported_files.append(file)

# Report unsupported files
if unsupported_files:

   print("\nUnsupported files detected:\n")

   for file in unsupported_files:

       print(
           f"Skipping: {file} "
           f"(type {Path(file).suffix} is not supported)"
       )

   print()

if not supported_files:
   raise FileNotFoundError(
       "No supported files found in the current directory."
   )

supported_files = [
   file for file in supported_files
   if not Path(file).name.startswith("masked_")
]

if not supported_files:
   raise FileNotFoundError(
       "No supported files found in the current directory."
   )
failed_files = []
successful_files = []

overall_start = time.perf_counter()

for file in supported_files:

   print(f"\nProcessing: {file}")

   file_start = time.perf_counter()

   current_file_counts = {}

   # Same tags within each file, but separate tags between files
   current_value_map = {}
   current_type_counters = {}

   output_file = None

   try:

       # CSV Files
       if file.endswith(".csv"):

           df = pd.read_csv(
               file,
               dtype=str,
               keep_default_na=False
           )

           masked_df = df.copy()

           for column in masked_df.columns:
               masked_df[column] = (
                   masked_df[column]
                   .apply(lambda x: mask_pii(
                       x, current_file_counts, current_value_map, current_type_counters
                   ))
               )

           output_file = os.path.join(
               OUTPUT_DIR,
               f"masked_{file}"
           )

           masked_df.to_csv(
               output_file,
               index=False
           )

       # Excel Files
       elif file.endswith((".xlsx", ".xls")):

           sheets = pd.read_excel(
               file,
               sheet_name=None,
               dtype=str
           )

           # Always save as .xlsx
           output_file = os.path.join(
               OUTPUT_DIR,
               f"masked_{Path(file).stem}.xlsx"
           )

           with pd.ExcelWriter(
               output_file,
               engine="openpyxl"
           ) as writer:

               for sheet_name, df in sheets.items():

                   df = df.fillna("")

                   masked_df = df.copy()

                   for column in masked_df.columns:
                       masked_df[column] = (
                           masked_df[column]
                           .astype(str)
                           .apply(lambda x: mask_pii(
                               x, current_file_counts, current_value_map, current_type_counters
                           ))
                       )

                   masked_df.to_excel(
                       writer,
                       sheet_name=sheet_name,
                       index=False
                   )

       # JSON Files
       elif file.endswith(".json"):

           with open(file, "r", encoding="utf-8") as f:
               data = json.load(f)

           masked = mask_json(data, current_file_counts, current_value_map, current_type_counters)

           output_file = os.path.join(
               OUTPUT_DIR,
               f"masked_{file}"
           )

           with open(output_file, "w", encoding="utf-8") as f:
               json.dump(
                   masked,
                   f,
                   indent=4
               )

       # Text Files
       elif file.endswith(".txt"):

           with open(file, "r", encoding="utf-8") as f:
               text = f.read()

           masked = mask_pii(text, current_file_counts, current_value_map, current_type_counters)

           output_file = os.path.join(
               OUTPUT_DIR,
               f"masked_{file}"
           )

           with open(output_file, "w", encoding="utf-8") as f:
               f.write(masked)

       # Markdown Files
       elif file.endswith(".md"):

           with open(file, "r", encoding="utf-8") as f:
               text = f.read()

           masked = mask_pii(text, current_file_counts, current_value_map, current_type_counters)

           output_file = os.path.join(
               OUTPUT_DIR,
               f"masked_{file}"
           )

           with open(output_file, "w", encoding="utf-8") as f:
               f.write(masked)

       # PDF Files
       elif file.endswith(".pdf"):

           pdf = fitz.open(file)

           for page in pdf:

               # Extract text blocks with coordinates
               blocks = page.get_text("dict")["blocks"]

               for block in blocks:

                   if "lines" not in block:
                       continue

                   # Full-block PII detection across text spans
                   span_list = [
                       span
                       for line in block["lines"]
                       for span in line["spans"]
                   ]

                   if not span_list:
                       continue

                   block_text = "".join(span["text"] for span in span_list)

                   if not block_text.strip():
                       continue

                   results = detect_entities(block_text, current_file_counts)

                   if not results:
                       continue

                   span_spans = []
                   cursor = 0
                   for span in span_list:
                       span_spans.append((cursor, cursor + len(span["text"])))
                       cursor += len(span["text"])

                   for span, (span_start, span_end) in zip(span_list, span_spans):
                       touched = any(
                           r.start < span_end and r.end > span_start
                           for r in results
                       )
                       if touched:
                           rect = fitz.Rect(span["bbox"])
                           page.add_redact_annot(
                               rect,
                               fill=(0, 0, 0)
                           )

               # Apply redactions
               page.apply_redactions()

           output_file = os.path.join(
               OUTPUT_DIR,
               f"masked_{Path(file).stem}.pdf"
           )

           pdf.save(output_file)

           pdf.close()

       # Word Documents
       elif file.endswith(".docx"):

           doc = Document(file)

           for paragraph in doc.paragraphs:
               mask_paragraph(paragraph, current_file_counts, current_value_map, current_type_counters)

           for table in doc.tables:
               mask_table(table, current_file_counts, current_value_map, current_type_counters)

           for section in doc.sections:
               regions = (
                   section.header, section.footer,
                   section.first_page_header, section.first_page_footer,
                   section.even_page_header, section.even_page_footer,
               )
               for region in regions:
                   for paragraph in region.paragraphs:
                       mask_paragraph(paragraph, current_file_counts, current_value_map, current_type_counters)
                   for table in region.tables:
                       mask_table(table, current_file_counts, current_value_map, current_type_counters)

           output_file = os.path.join(
               OUTPUT_DIR,
               f"masked_{file}"
           )

           doc.save(output_file)


       # XML Files
       elif file.endswith(".xml"):

           tree = ET.parse(file)

           root = tree.getroot()

           def mask_xml(element):
               if element.text:
                   element.text = mask_pii(
                       element.text, current_file_counts, current_value_map, current_type_counters
                   )
               if element.tail:
                   element.tail = mask_pii(
                       element.tail, current_file_counts, current_value_map, current_type_counters
                   )
               for attr_name, attr_value in element.attrib.items():
                   element.set(
                       attr_name,
                       mask_pii(attr_value, current_file_counts, current_value_map, current_type_counters)
                   )
               for child in element:
                   mask_xml(child)

           mask_xml(root)

           output_file = os.path.join(
               OUTPUT_DIR,
               f"masked_{file}"
           )

           tree.write(
               output_file,
               encoding="utf-8",
               xml_declaration=True
           )

       else:
           print(f"Unsupported file type: {file}")
           continue

       print(f"Saved: {output_file}")

       file_pii_counts[file] = current_file_counts

       risk_score = sum(current_file_counts.values())

       if (
           current_file_counts.get("US_SSN", 0)
           or current_file_counts.get("CREDIT_CARD", 0)
       ):
           risk = "HIGH"

       elif risk_score >= 30:
           risk = "HIGH"

       elif risk_score >= 10:
           risk = "MEDIUM"

       else:
           risk = "LOW"

       file_risk[file] = risk

       print(f"Risk Level: {risk}")

       file_time = time.perf_counter() - file_start
       print(f"Processing Time: {file_time:.2f} seconds")

       successful_files.append(file)

   except Exception as e:

       print(f"\nERROR processing: {file}")
       print(f"Reason: {e}")
       print("Skipping this file and continuing...\n")

       failed_files.append(file)

       continue
   
# Raw PII Detection Report
print("\nFinished processing files.\n")

print("\nRaw PII Detection Report:\n")

total_pii = sum(pii_counts.values())

sorted_entities = sorted(pii_counts.items(), key=lambda x: x[1], reverse=True)

for entity, count in sorted_entities:
  percentage = (count / total_pii * 100) if total_pii > 0 else 0
  print(f"{entity}: {count} ({percentage:.1f}%)")

print("\nTOTAL PII MASKED:")
print(total_pii)

print("\nPII Found by File:\n")
for filename, counts in file_pii_counts.items():
    print(filename)
    print("-" * len(filename))
    if counts:
        for entity, total in counts.items():
            print(f"{entity}: {total}")
    else:
        print("No PII Detected")
    print()

print("\nRisk Summary\n")
for file, risk in file_risk.items():
    print(f"{file}: {risk}")

overall_time = time.perf_counter() - overall_start
print(
    f"\nTotal Runtime: {overall_time:.2f} seconds"
)

print("\nFile Processing Summary:")
print(f"Successful files: {len(successful_files)}")
print(f"Unsupported files: {len(unsupported_files)}")
print(f"Failed files: {len(failed_files)}")

log_file = os.path.join(
    OUTPUT_DIR,
    "masking_log.txt"
)

with open(log_file, "w", encoding="utf-8") as log:

    log.write("RISBDC PII Masking Report\n")
    log.write("=" * 40 + "\n\n")

    log.write(
        f"Generated: "
        f"{datetime.now()}\n\n"
    )

    log.write(
        f"Filter Preset: "
        f"{args.preset}\n\n"
    )

    log.write(
        f"Successful Files: "
        f"{len(successful_files)}\n"
    )

    log.write(
        f"Unsupported Files: "
        f"{len(unsupported_files)}\n"
    )

    log.write(
        f"Failed Files: "
        f"{len(failed_files)}\n\n"
    )

    log.write(
        f"Total Runtime: "
        f"{overall_time:.2f} seconds\n\n"
    )

    log.write(
        f"Total PII Masked: "
        f"{total_pii}\n\n"
    )

    log.write("Per File Summary\n")
    log.write("-" * 25 + "\n")

    for filename, counts in file_pii_counts.items():

        log.write(f"\n{filename}\n")

        for entity, total in counts.items():

            log.write(
                f"   {entity}: {total}\n"
            )

        log.write(
            f"Risk Level: "
            f"{file_risk[filename]}\n"
        )

print(f"\nMasking log saved to: {log_file}")